"""
Tests for the list-view quick export (horilla_views/generic/cbv/views.py,
HorillaListView.dispatch and export_data, generic/export_fields_modal.html).

The expected outputs below were captured from the export code as it was
before the exec() removal, so they prove that legitimate exports keep the
same bytes (CSV, JSON) and cells (XLSX). No output contains row ids, so the
goldens do not depend on database sequences and hold under --keepdb.

The other tests cover the security fix: nothing from the request is compiled
or executed, only the view's own columns are exported, the export goes through
the list's own login and permission decorators, and rendering a list no longer
adds URL patterns.

Synthetic data only. The test runner never touches the real database.
"""

import builtins
import io
import json
import threading
from unittest import mock

from django.contrib.auth.models import AnonymousUser, Group, Permission, User
from django.contrib.sessions.backends.db import SessionStore
from django.db import connection
from django.test import Client, RequestFactory, TestCase, override_settings
from django.urls import get_resolver, reverse
from openpyxl import load_workbook

from employee.models import Employee
from horilla.horilla_middlewares import _thread_locals
from horilla_views.generic.cbv.views import HorillaListView
from leave.tests import add_missing_user_column
from pms.cbvs import BonusPointSettingListView
from pms.models import BonusPointSetting

PASSWORD = "not-a-real-password"
COLUMNS = [["Name", "name"], ["Key", "natural_key"]]
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

GOLDEN_CSV = "".join(
    [
        "Name,Key\r\n",
        "Alpha,\"('Alpha',)\"\r\n",
        "Beta   Team,\"('Beta   Team',)\"\r\n",
    ]
)
GOLDEN_JSON = (
    "[\n"
    "    {\n"
    '        "Name": "Alpha",\n'
    '        "Key": "(\'Alpha\',)"\n'
    "    },\n"
    "    {\n"
    '        "Name": "Beta   Team",\n'
    '        "Key": "(\'Beta   Team\',)"\n'
    "    }\n"
    "]"
)
GOLDEN_XLSX_ROWS = [
    ["Name", "Key"],
    ["Alpha", "('Alpha',)"],
    ["Beta   Team", "('Beta   Team',)"],
]
GOLDEN_PMS_CSV = (
    "Model,Applicable For,Bonus For,Condition,Points,Is Active\r\n"
    "Objective,Owner,Completing,Completion Date < End Date,5,\r\n"
)


class GroupExportList(HorillaListView):
    """Test-only list over auth.Group: one model field and one method column."""

    model = Group
    columns = [("Name", "name"), ("Key", "natural_key")]
    export_fields = []


def make_login(label, superuser=False):
    """A login with its Employee, as every real user has; menus read it."""
    create = User.objects.create_superuser if superuser else User.objects.create_user
    user = create(username=label, password=PASSWORD)
    employee = Employee.objects.create(
        employee_user_id=user,
        employee_first_name=label,
        email=f"{label}@example.invalid",
        phone="0700000000",
    )
    return user, employee


def xlsx_rows(response):
    """Sheet title and cell values; the file bytes carry a fresh timestamp."""
    sheet = load_workbook(io.BytesIO(response.content)).active
    return sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)]


class ExportTestCase(TestCase):
    @classmethod
    def setUpClass(cls):
        add_missing_user_column()
        super().setUpClass()

    def setUp(self):
        # ThreadLocalMiddleware never clears the request; reset it per test.
        _thread_locals.request = None
        self.addCleanup(setattr, _thread_locals, "request", None)
        self.user, _ = make_login("exporter")
        self.alpha = Group.objects.create(name="Alpha")
        self.beta = Group.objects.create(name="Beta   Team")
        self.ids = [self.alpha.pk, self.beta.pk]

    def export_request(self, data, user=None):
        """A quick-export POST as the modal sends it."""
        request = RequestFactory().post("/test-export/", data)
        request.user = user or self.user
        # Display methods such as is_active_toggle render templates whose
        # context processors read the session key.
        request.session = SessionStore()
        request.session.create()
        return request

    def build_view(self, request, view_cls=GroupExportList):
        # HorillaListView.__init__ reads the request from the thread-local and
        # filters ToggleColumn by its user, which fails for an AnonymousUser.
        if request.user.is_authenticated:
            _thread_locals.request = request
        view = view_cls()
        _thread_locals.request = request
        view.request = request
        return view

    def export(self, columns, ids=None, fmt="csv", view_cls=GroupExportList):
        """Call export_data directly with a well-formed POST."""
        request = self.export_request(
            {
                "ids": json.dumps(self.ids if ids is None else ids),
                "columns": json.dumps(columns),
                "format": fmt,
            }
        )
        return self.build_view(request, view_cls).export_data(request)


class ExportOutputUnchangedTests(ExportTestCase):
    """Legitimate exports give the same output as before the fix."""

    def test_csv_output_unchanged(self):
        response = self.export(COLUMNS, fmt="csv")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "text/csv")
        self.assertEqual(
            response["Content-Disposition"], 'attachment; filename="quick_export.csv"'
        )
        self.assertEqual(response.content.decode(), GOLDEN_CSV)

    def test_json_output_unchanged(self):
        response = self.export(COLUMNS, fmt="json")
        self.assertEqual(response["Content-Type"], "application/json")
        self.assertEqual(
            response["Content-Disposition"], 'attachment; filename="quick_export.json"'
        )
        self.assertEqual(response.content.decode(), GOLDEN_JSON)

    def test_xlsx_cells_unchanged(self):
        response = self.export(COLUMNS, fmt="xlsx")
        self.assertEqual(response["Content-Type"], XLSX_TYPE)
        self.assertEqual(
            response["Content-Disposition"], 'attachment; filename="quick_export.xlsx"'
        )
        self.assertEqual(xlsx_rows(response), ("Quick Export", GOLDEN_XLSX_ROWS))

    def test_pdf_still_generated(self):
        response = self.export(COLUMNS, fmt="pdf")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertEqual(
            response["Content-Disposition"], 'attachment; filename="quick_export.pdf"'
        )
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_unknown_format_falls_back_to_xlsx(self):
        response = self.export(COLUMNS, fmt="xml")
        self.assertEqual(response["Content-Type"], XLSX_TYPE)
        self.assertEqual(xlsx_rows(response), ("Quick Export", GOLDEN_XLSX_ROWS))

    def test_duplicate_and_reordered_columns_unchanged(self):
        columns = [["Key", "natural_key"], ["Name", "name"], ["Name", "name"]]
        response = self.export(columns)
        self.assertEqual(
            response.content.decode(),
            "Key,Name\r\n"
            "\"('Alpha',)\",Alpha\r\n"
            "\"('Beta   Team',)\",Beta   Team\r\n",
        )

    def test_unicode_label_unchanged(self):
        response = self.export([["Jina ✓ <b>", "name"]])
        self.assertEqual(
            response.content.decode(), "Jina ✓ <b>\r\nAlpha\r\nBeta   Team\r\n"
        )

    def test_empty_ids_exports_headers_only(self):
        response = self.export(COLUMNS, ids=[])
        self.assertEqual(response.content.decode(), "Name,Key\r\n")

    def test_export_twice_identical(self):
        from horilla.urls import urlpatterns

        before = len(urlpatterns)
        first = self.export(COLUMNS).content
        second = self.export(COLUMNS).content
        self.assertEqual(first, second)
        self.assertEqual(len(urlpatterns), before)

    def test_pms_export_direct_call_unchanged(self):
        # Covers the <input> branch of remove_extra_spaces (is_active_toggle).
        setting = BonusPointSetting.objects.create(
            model="pms.models.EmployeeObjective",
            applicable_for="owner",
            bonus_for="completed",
            field_1="complition_date",
            conditions="<",
            field_2="end_date",
            points=5,
            is_active=True,
        )
        columns = [list(column) for column in BonusPointSettingListView.columns]
        response = self.export(
            columns, ids=[setting.pk], view_cls=BonusPointSettingListView
        )
        self.assertEqual(response.content.decode(), GOLDEN_PMS_CSV)


# A column name that, before the fix, ran its second line as Python while the
# export resource class was being built. The trailing # comments out the rest
# of the generated source, so the code compiled and ran.
PAYLOAD = "a(self,i):return 1\nimport builtins;builtins.PWNED=1#"
VIEWS_LOGGER = "horilla_views.generic.cbv.views"
EXPORT_FLAG = "export_list_view"  # hidden field in generic/export_fields_modal.html
PMS_COLUMNS = [list(column) for column in BonusPointSettingListView.columns]


class NoExportGroupList(GroupExportList):
    quick_export = False


class NoExportBonusPointList(BonusPointSettingListView):
    quick_export = False


def all_subclasses(cls):
    for subclass in cls.__subclasses__():
        yield subclass
        yield from all_subclasses(subclass)


class ExportInputTests(ExportTestCase):
    """Only the view's own columns are exported; nothing is executed."""

    def setUp(self):
        super().setUp()
        self.addCleanup(lambda: builtins.__dict__.pop("PWNED", None))

    def test_exec_payload_in_column_name_is_not_executed(self):
        with self.assertLogs(VIEWS_LOGGER, "WARNING") as logs:
            response = self.export([["x", PAYLOAD], ["Name", "name"]])
        self.assertFalse(hasattr(builtins, "PWNED"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content.decode(), "Name\r\nAlpha\r\nBeta   Team\r\n")
        self.assertEqual(
            logs.output,
            [
                f"WARNING:{VIEWS_LOGGER}:List export ignored 1 unknown column(s) "
                "on GroupExportList"
            ],
        )

    def test_exec_payload_with_empty_ids_is_not_executed(self):
        # The old code ran the payload while building the class, with no rows.
        with self.assertLogs(VIEWS_LOGGER, "WARNING"):
            response = self.export([["x", PAYLOAD]], ids=[])
        self.assertFalse(hasattr(builtins, "PWNED"))
        self.assertEqual(response.status_code, 200)

    def test_method_column_outside_allow_list_is_not_called(self):
        # getattribute calls zero-argument methods, so "delete" deleted rows.
        with self.assertLogs(VIEWS_LOGGER, "WARNING"):
            response = self.export(
                [["Gone", "delete"], ["Gone too", "natural_key__delete"]],
                ids=[self.alpha.pk],
            )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(Group.objects.filter(pk=self.alpha.pk).exists())
        self.assertEqual(response.content.decode(), "\r\n")

    def test_attribute_path_outside_allow_list_is_dropped(self):
        self.user.groups.add(self.alpha)
        with self.assertLogs(VIEWS_LOGGER, "WARNING"):
            response = self.export(
                [["Name", "name"], ["Hash", "user_set__password"]],
                ids=[self.alpha.pk],
            )
        self.assertNotIn(self.user.password, response.content.decode())
        self.assertEqual(response.content.decode(), "Name\r\nAlpha\r\n")

    def test_unknown_columns_dropped_and_logged_without_values(self):
        columns = [
            ["S", "s"],
            ["AB", "a-b"],
            ["Five", 5],
            ["List", ["name"]],
            ["Naive", "naïve"],
            [5, "name"],
            ["Name", "name"],
        ]
        with self.assertLogs(VIEWS_LOGGER, "WARNING") as logs:
            response = self.export(columns)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content.decode(), "Name\r\nAlpha\r\nBeta   Team\r\n")
        self.assertEqual(
            logs.output,
            [
                f"WARNING:{VIEWS_LOGGER}:List export ignored 6 unknown column(s) "
                "on GroupExportList"
            ],
        )

    def test_malformed_input_returns_400(self):
        good_ids, good_columns = json.dumps(self.ids), json.dumps(COLUMNS)
        deep = "[" * 300 + "]" * 300
        cases = {
            "no ids or columns": {},
            "no columns": {"ids": good_ids},
            "no ids": {"columns": good_columns},
            "columns not a literal": {"ids": good_ids, "columns": "not a list"},
            "three-item column": {"ids": good_ids, "columns": '[["a", "b", "c"]]'},
            "columns of ints": {"ids": good_ids, "columns": "[1, 2]"},
            "columns a string": {"ids": good_ids, "columns": "'ab'"},
            "non-numeric id": {"ids": "['abc']", "columns": good_columns},
            "ids not a list": {"ids": "5", "columns": good_columns},
            "deeply nested ids": {"ids": deep, "columns": good_columns},
            "deeply nested columns": {"ids": good_ids, "columns": deep},
        }
        for name, data in cases.items():
            with self.subTest(name):
                request = self.export_request(data)
                response = self.build_view(request).export_data(request)
                self.assertEqual(response.status_code, 400)
                self.assertEqual(response.content, b"Invalid export request")

    def test_parallel_exports_independent(self):
        # ids=[] because a new thread has its own database connection and
        # cannot see this test's uncommitted rows; id__in=[] runs no query.
        requests = {
            "Name": self.export_request(
                {
                    "ids": "[]",
                    "columns": json.dumps([["Name", "name"]]),
                    "format": "csv",
                }
            ),
            "Key": self.export_request(
                {
                    "ids": "[]",
                    "columns": json.dumps([["Key", "natural_key"]]),
                    "format": "csv",
                }
            ),
        }
        barrier = threading.Barrier(len(requests))
        results = {}

        def run(label, request):
            try:
                _thread_locals.request = request
                view = GroupExportList()
                view.request = request
                barrier.wait(timeout=10)
                results[label] = view.export_data(request).content.decode()
            except Exception as error:  # surfaced by the assertEqual below
                results[label] = repr(error)
            finally:
                connection.close()

        threads = [threading.Thread(target=run, args=item) for item in requests.items()]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(timeout=30)
        self.assertEqual(results, {"Name": "Name\r\n", "Key": "Key\r\n"})


class ExportDispatchTests(ExportTestCase):
    """The export is served by dispatch() and checks who is asking."""

    def flagged_request(self, user=None):
        return self.export_request(
            {
                EXPORT_FLAG: "1",
                "ids": json.dumps(self.ids),
                "columns": json.dumps(COLUMNS),
                "format": "csv",
            },
            user=user,
        )

    def test_flagged_post_is_served_as_export(self):
        request = self.flagged_request()
        response = self.build_view(request).dispatch(request)
        self.assertEqual(response.content.decode(), GOLDEN_CSV)

    def test_anonymous_export_denied(self):
        # A view with no decorators, like pms EmployeeBonusPointListView.
        request = self.flagged_request(user=AnonymousUser())
        response = self.build_view(request).dispatch(request)
        self.assertEqual(response.status_code, 403)
        self.assertNotIn("Content-Disposition", response)

    def test_quick_export_disabled_denied(self):
        request = self.flagged_request()
        response = self.build_view(request, NoExportGroupList).dispatch(request)
        self.assertEqual(response.status_code, 403)
        self.assertNotIn("Content-Disposition", response)

    def test_all_list_view_columns_are_identifiers(self):
        # Import every module that defines a list view.
        get_resolver().url_patterns
        checked = 0
        for cls in set(all_subclasses(HorillaListView)):
            if "<locals>" in cls.__qualname__:
                continue
            # Some views declare columns as a cached_property on the instance.
            columns = cls.__new__(cls).columns
            for column in columns:
                name = column if isinstance(column, str) else column[1]
                with self.subTest(view=cls.__qualname__, column=name):
                    self.assertTrue(name.isidentifier())
                    self.assertFalse(name.endswith(("delete", "save")))
            checked += 1
        self.assertGreater(checked, 10)


class ListExportAccessTests(ExportTestCase):
    """End to end through the real PMS bonus point settings list."""

    def setUp(self):
        super().setUp()
        self.url = reverse("bonus-point-setting-list-view")
        self.hr_user, hr_employee = make_login("hr")
        self.hr_user.user_permissions.add(
            Permission.objects.get(
                codename="view_bonuspointsetting", content_type__app_label="pms"
            )
        )
        # The plain employee is also a reporting manager (manager scope is p3-02).
        self.plain_user, plain_employee = make_login("plain")
        work_info = hr_employee.employee_work_info
        work_info.reporting_manager_id = plain_employee
        work_info.save()
        self.superuser, _ = make_login("root", superuser=True)
        self.setting = BonusPointSetting.objects.create(
            model="pms.models.EmployeeObjective",
            applicable_for="owner",
            bonus_for="completed",
            field_1="complition_date",
            conditions="<",
            field_2="end_date",
            points=5,
            is_active=True,
        )

    def login_client(self, user=None):
        client = Client(raise_request_exception=False)
        if user:
            client.force_login(user)
        return client

    def post_export(self, user=None, **headers):
        return self.login_client(user).post(
            self.url,
            {
                EXPORT_FLAG: "1",
                "ids": json.dumps([self.setting.pk]),
                "columns": json.dumps(PMS_COLUMNS),
                "format": "csv",
            },
            **headers,
        )

    def assertNoFile(self, response):
        self.assertNotIn("attachment", response.get("Content-Disposition", ""))
        self.assertNotIn(b"Objective", response.content)

    def export_routes(self):
        from horilla.urls import urlpatterns

        return [p for p in urlpatterns if "export-list-view" in str(p.pattern)]

    def test_list_render_adds_no_export_route(self):
        client = self.login_client(self.hr_user)
        before = len(self.export_routes())
        for _ in range(3):
            response = client.get(self.url, HTTP_HX_REQUEST="true")
            self.assertEqual(response.status_code, 200)
        self.assertEqual(len(self.export_routes()), before)

    def test_modal_posts_to_list_url(self):
        response = self.login_client(self.hr_user).get(
            f"{self.url}?search=Obj", HTTP_HX_REQUEST="true"
        )
        self.assertContains(response, f'action="{self.url}?search=Obj"')
        self.assertContains(
            response, f'<input type="hidden" name="{EXPORT_FLAG}" value="1">'
        )

    def test_modal_action_keeps_encoded_query(self):
        response = self.login_client(self.hr_user).get(
            f"{self.url}?search=Jina%20%C3%A9&page=2", HTTP_HX_REQUEST="true"
        )
        self.assertContains(
            response, f'action="{self.url}?search=Jina%20%C3%A9&amp;page=2"'
        )

    def test_pms_export_anonymous_gets_no_file(self):
        response = self.post_export()
        self.assertNotEqual(response.status_code, 200)
        self.assertNoFile(response)

    def test_pms_export_without_permission_gets_no_file(self):
        response = self.post_export(self.plain_user)
        self.assertNoFile(response)

    def test_pms_export_with_permission_unchanged(self):
        response = self.post_export(self.hr_user)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Disposition"], 'attachment; filename="quick_export.csv"'
        )
        self.assertEqual(response.content.decode(), GOLDEN_PMS_CSV)

    def test_pms_export_superuser(self):
        response = self.post_export(self.superuser)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content.decode(), GOLDEN_PMS_CSV)

    def test_export_with_hx_header_served(self):
        response = self.post_export(self.hr_user, HTTP_HX_REQUEST="true")
        self.assertEqual(response.content.decode(), GOLDEN_PMS_CSV)

    def test_non_export_post_without_hx_still_405(self):
        response = self.login_client(self.hr_user).post(
            self.url,
            {"ids": json.dumps([self.setting.pk]), "columns": json.dumps(PMS_COLUMNS)},
        )
        self.assertTemplateUsed(response, "405.html")
        self.assertNoFile(response)

    def test_export_flag_in_query_string_ignored(self):
        client = self.login_client(self.hr_user)
        response = client.get(f"{self.url}?{EXPORT_FLAG}=1")
        self.assertTemplateUsed(response, "405.html")
        response = client.get(f"{self.url}?{EXPORT_FLAG}=1", HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        self.assertNotIn("Content-Disposition", response)

    def test_old_export_path_404(self):
        response = self.login_client(self.hr_user).post("/export-list-view-hlv0000/")
        self.assertEqual(response.status_code, 404)

    def test_denied_export_under_horilla_login_required_is_403(self):
        # horilla's login_required turns exceptions into went_wrong.html, or
        # runs the view a second time when DEBUG is on. Denials must be plain
        # responses, the same in both modes, and run the export once.
        cases = {
            "export disabled": (NoExportBonusPointList, json.dumps(PMS_COLUMNS), 403),
            "malformed columns": (BonusPointSettingListView, "not a list", 400),
        }
        for debug in (True, False):
            for name, (view_cls, columns, status) in cases.items():
                with self.subTest(name, debug=debug), override_settings(DEBUG=debug):
                    request = RequestFactory().post(
                        self.url,
                        {
                            EXPORT_FLAG: "1",
                            "ids": json.dumps([self.setting.pk]),
                            "columns": columns,
                        },
                    )
                    request.user = self.hr_user
                    request.session = SessionStore()
                    _thread_locals.request = request
                    with mock.patch.object(
                        view_cls,
                        "export_data",
                        autospec=True,
                        side_effect=HorillaListView.export_data,
                    ) as export_data:
                        response = view_cls.as_view()(request)
                    self.assertEqual(response.status_code, status)
                    self.assertEqual(export_data.call_count, 1)
